from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[4]
REPORT_DIR = ROOT / "Sprint 2" / "reports" / "progress"
ASSET_DIR = ROOT / "Sprint 2" / "assets" / "report-captures"
OUTPUT = REPORT_DIR / "Sprint_2_W1_Progress_Report_Inteli_YvY_Capital.docx"

PAGE_TOTAL = 20
GENERATION_DATE = "August 22, 2026"

FONT = "Segoe UI"
FONT_BOLD = "Segoe UI Semibold"

NAVY = "2B203D"
NAVY_2 = "392653"
PURPLE = "613F7B"
LAVENDER = "EEE8F2"
LAVENDER_2 = "F7F3F8"
CORAL = "FF5364"
TEAL = "2D9291"
TEAL_LIGHT = "E4F2F0"
ORANGE = "E49324"
ORANGE_LIGHT = "FAEED8"
RED = "B33A54"
RED_LIGHT = "F8E3E9"
INK = "292431"
MUTED = "716B78"
LINE = "D8D0DC"
WHITE = "FFFFFF"
CREAM = "FCFAF7"

STATUS_COLORS = {
    "Implemented and visible": (TEAL_LIGHT, "215F58"),
    "Interactive prototype": (LAVENDER, PURPLE),
    "Partially implemented": (ORANGE_LIGHT, "8D5A0A"),
    "Specified but not yet visible": ("ECEAF0", "5E5867"),
    "Planned for a later sprint": ("E7EFF7", "365D7D"),
    "Pending partner validation": (ORANGE_LIGHT, "8D5A0A"),
    "Out of scope for Sprint 2 Week 1": ("F1F0F2", "66616C"),
}


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=120, bottom=55, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            set_cell_width(cell, width)
            set_cell_margins(cell)


def set_table_borders(table, color=LINE, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def remove_table_borders(table):
    set_table_borders(table, WHITE, 0)
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    for edge in borders:
        edge.set(qn("w:val"), "nil")


def set_run(run, size=9.2, color=INK, bold=False, italic=False, font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def style_paragraph(paragraph, before=0, after=4, line=1.08, align=None, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if keep:
        fmt.keep_with_next = True


def add_text(container, text, size=9.2, color=INK, bold=False, italic=False, before=0, after=4, line=1.08, align=None, keep=False):
    p = container.add_paragraph()
    style_paragraph(p, before, after, line, align, keep)
    set_run(p.add_run(text), size, color, bold, italic, FONT_BOLD if bold else FONT)
    return p


def add_rich_text(container, parts, before=0, after=4, line=1.08, align=None, keep=False):
    p = container.add_paragraph()
    style_paragraph(p, before, after, line, align, keep)
    for text, options in parts:
        set_run(p.add_run(text), **options)
    return p


def add_bullet(container, text, level=0, size=8.8, color=INK, after=2.5):
    p = container.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, 0, after, 1.06)
    set_run(p.add_run(text), size, color)
    return p


def add_number(container, text, size=8.8, after=2.5):
    p = container.add_paragraph(style="List Number")
    style_paragraph(p, 0, after, 1.06)
    set_run(p.add_run(text), size, INK)
    return p


def add_kicker(container, text, color=PURPLE, after=3):
    return add_text(container, text.upper(), 8.3, color, True, after=after, keep=True)


def add_heading(container, text, level=1, after=None, before=None):
    p = container.add_paragraph(style=f"Heading {level}")
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), 20 if level == 1 else 13.5 if level == 2 else 10.5, NAVY, True, font=FONT_BOLD)
    return p


def add_callout(container, title, body, fill=NAVY, accent=CORAL, text_color=WHITE, compact=False):
    table = container.add_table(rows=1, cols=1)
    set_table_geometry(table, [10000], 0)
    remove_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    margin = 120 if compact else 180
    set_cell_margins(cell, top=margin, start=220, bottom=margin, end=220)
    add_text(cell, title.upper(), 8.0 if compact else 8.2, accent, True, after=3 if compact else 5)
    add_text(cell, body, 9.4 if compact else 10.2, text_color, True, after=0, line=1.05 if compact else 1.12)
    add_text(container, "", 1, after=2 if compact else 4)
    return table


def add_metric_grid(container, metrics, compact=False):
    table = container.add_table(rows=2, cols=2)
    set_table_geometry(table, [5000, 5000], 0)
    set_table_borders(table, LINE, 4)
    for index, (value, label, detail, accent) in enumerate(metrics):
        cell = table.cell(index // 2, index % 2)
        shade_cell(cell, LAVENDER_2)
        margin = 95 if compact else 150
        set_cell_margins(cell, margin, 170, margin, 170)
        add_text(cell, value, 16 if compact else 18, accent, True, after=1)
        add_text(cell, label, 8.8 if compact else 9.4, NAVY, True, after=1 if compact else 2)
        add_text(cell, detail, 7.2 if compact else 7.8, MUTED, after=0)
    add_text(container, "", 1, after=1 if compact else 3)
    return table


def add_fact_table(container, rows, label_width=2300, total_width=10000, font_size=8.2):
    table = container.add_table(rows=0, cols=2)
    set_table_geometry(table, [label_width, total_width - label_width], 0)
    set_table_borders(table, LINE, 4)
    for label, value in rows:
        cells = table.add_row().cells
        shade_cell(cells[0], LAVENDER_2)
        add_text(cells[0], label, font_size, PURPLE, True, after=0)
        add_text(cells[1], value, font_size, INK, after=0, line=1.04)
    return table


def add_status(container, status, detail=""):
    fill, text_color = STATUS_COLORS[status]
    table = container.add_table(rows=1, cols=1)
    set_table_geometry(table, [10000], 0)
    remove_table_borders(table)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_margins(cell, 90, 140, 90, 140)
    add_rich_text(
        cell,
        [
            (status, {"size": 8.4, "color": text_color, "bold": True}),
            ((f" - {detail}" if detail else ""), {"size": 8.4, "color": text_color}),
        ],
        after=0,
    )
    return table


def add_matrix(container, headers, rows, widths, font_size=7.0, status_col=None):
    table = container.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, 0)
    set_table_borders(table, LINE, 4)
    hdr = table.rows[0].cells
    for idx, title in enumerate(headers):
        shade_cell(hdr[idx], NAVY)
        add_text(hdr[idx], title, 7.1, WHITE, True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if row_idx % 2:
                shade_cell(cells[idx], LAVENDER_2)
            if status_col is not None and idx == status_col and value in STATUS_COLORS:
                fill, text_color = STATUS_COLORS[value]
                shade_cell(cells[idx], fill)
                add_text(cells[idx], value, font_size, text_color, True, after=0, line=1.0)
            else:
                add_text(cells[idx], value, font_size, INK, idx == 0, after=0, line=1.0)
    return table


def add_screenshot(container, filename, caption, width_inches=6.95):
    path = ASSET_DIR / filename
    p = container.add_paragraph()
    style_paragraph(p, 2, 2, 1, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    add_text(container, caption, 7.6, MUTED, italic=True, after=5, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)


def set_page_geometry(section, landscape=False, cover=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    if cover:
        section.top_margin = Cm(0)
        section.bottom_margin = Cm(0)
        section.left_margin = Cm(0)
        section.right_margin = Cm(0)
        section.header_distance = Cm(0)
        section.footer_distance = Cm(0)
    else:
        section.top_margin = Cm(1.35)
        section.bottom_margin = Cm(1.45)
        section.left_margin = Cm(1.65)
        section.right_margin = Cm(1.65)
        section.header_distance = Cm(0.48)
        section.footer_distance = Cm(0.38)


def configure_header_footer(section, page_number, page_label, landscape=False):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    footer = section.footer
    for p in header.paragraphs:
        p._element.getparent().remove(p._element)
    for p in footer.paragraphs:
        p._element.getparent().remove(p._element)
    width = 14400 if landscape else 10000
    header_table = header.add_table(rows=1, cols=2, width=Inches(10 if landscape else 6.95))
    set_table_geometry(header_table, [width // 2, width // 2], 0)
    remove_table_borders(header_table)
    add_text(header_table.cell(0, 0), "INTELI × YVY CAPITAL", 7.7, NAVY, True, after=0)
    add_text(header_table.cell(0, 1), page_label.upper(), 7.2, MUTED, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    rule = header.add_paragraph()
    style_paragraph(rule, 0, 0, 1)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:color"), CORAL)
    bottom.set(qn("w:space"), "1")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    footer_table = footer.add_table(rows=1, cols=3, width=Inches(10 if landscape else 6.95))
    set_table_geometry(footer_table, [2300 if not landscape else 3300, 5600 if not landscape else 8000, 2100 if not landscape else 2700], 0)
    remove_table_borders(footer_table)
    for cell in footer_table.rows[0].cells:
        shade_cell(cell, NAVY)
        set_cell_margins(cell, 80, 130, 80, 130)
    add_text(footer_table.cell(0, 0), "inteli | Special Projects", 7.0, WHITE, True, after=0)
    add_text(footer_table.cell(0, 1), "Funds Manager - Working Draft", 6.8, "D8D1DE", after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(footer_table.cell(0, 2), f"{page_number:02d}/{PAGE_TOTAL:02d}", 7.0, WHITE, True, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)


def new_page(doc, page_number, label, title, subtitle=None, landscape=False):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(section, landscape=landscape)
    configure_header_footer(section, page_number, label, landscape)
    add_kicker(doc, label, PURPLE, after=2)
    p = doc.add_paragraph()
    style_paragraph(p, 0, 3, 1.0, keep=True)
    set_run(p.add_run(title), 21 if not landscape else 17.5, NAVY, True, font=FONT_BOLD)
    if subtitle:
        add_text(doc, subtitle, 9.0 if not landscape else 8.3, MUTED, after=6 if not landscape else 4, line=1.05)
    else:
        add_text(doc, "", 1, after=3)
    return section


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (("Heading 1", 20, 0, 5), ("Heading 2", 13.5, 8, 4), ("Heading 3", 10.5, 6, 3)):
        style = styles[name]
        style.font.name = FONT_BOLD
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_BOLD)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_BOLD)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(8.8)
        style.paragraph_format.space_after = Pt(2.5)
        style.paragraph_format.line_spacing = 1.06


def add_cover(doc):
    section = doc.sections[0]
    set_page_geometry(section, cover=True)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [11906], 0)
    remove_table_borders(table)
    row = table.rows[0]
    row.height = Cm(29.68)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade_cell(cell, NAVY)
    set_cell_margins(cell, 900, 1050, 700, 1050)
    add_text(cell, "INTELI × YVY CAPITAL", 13, CORAL, True, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "FUNDS MANAGER", 33, WHITE, True, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "SPRINT 2 — WEEK 1", 12, CORAL, True, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "PROGRESS REPORT", 24, WHITE, True, after=9, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "Dashboard Prototype, Current Capabilities and Path to the Final Product", 12, "E6DFEA", after=28, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    status = cell.add_table(rows=1, cols=1)
    set_table_geometry(status, [5600], 0)
    status.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(status, CORAL, 7)
    shade_cell(status.cell(0, 0), NAVY_2)
    add_text(status.cell(0, 0), "Working Draft — Pending Partner Review", 9.4, WHITE, True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "", 1, after=22)
    add_text(cell, GENERATION_DATE, 10, WHITE, True, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "Project deliverable prepared by the Inteli student team in partnership with YvY Capital.", 8.4, "D0C8D7", after=28, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(cell, "Prototype evidence only | Synthetic and aliased data | Not a production release", 7.7, CORAL, True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)


def build_report():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    props = doc.core_properties
    props.title = "Sprint 2 - Week 1 Progress Report"
    props.subject = "Funds Manager dashboard prototype progress and final-product alignment"
    props.author = "Inteli student team in partnership with YvY Capital"
    props.keywords = "Funds Manager, Sprint 2, dashboard prototype, data governance"

    add_cover(doc)

    # Page 2 - Sections 2 and 3
    new_page(doc, 2, "Executive Summary", "Executive Summary", "Verified progress, current boundaries, and the decisions needed to move from prototype to product.")
    add_callout(doc, "Current position", "Sprint 2 Week 1 has produced a coherent English-language interactive prototype that demonstrates the target information hierarchy and two priority journeys. It is not a production application: the displayed values are synthetic, role selection is simulated, and partner approval remains pending.")
    add_metric_grid(doc, [
        ("9", "Live dashboard sections", "Every section executes and was inspected.", PURPLE),
        ("10", "Demonstration scenarios", "Current, delayed, incomplete, empty, error, denied, and related states.", TEAL),
        ("2", "Priority user journeys", "Executive exploration and Analyst validation.", ORANGE),
        ("100%", "English UI coverage", "Visible and accessible copy checked in the DOM.", CORAL),
    ])
    add_heading(doc, "Sprint 2 Week 1 Objectives", 2)
    add_bullet(doc, "Translate priority user needs and the Sprint 1 roadmap into a demonstrable navigation and screen baseline.")
    add_bullet(doc, "Expose fund, performance, comparison, quality, import, and lineage concepts without presenting synthetic evidence as certified business data.")
    add_bullet(doc, "Make assumptions, freshness, quality, and unavailable prerequisites visible at the point of use.")
    add_bullet(doc, "Create evidence for partner review before production development and before the Sprint 2 exit gate is claimed.")
    add_status(doc, "Pending partner validation", "The prototype is a working draft; no YvY Capital approval is represented in this report.")
    add_text(doc, "Evidence basis: the verified dashboard, the project charter, the Sprint 1 delivery roadmap, the internal sprint plan, the Feature Guide, the Data Dictionary, and the Sprint 2 audit/remediation evidence.", 7.6, MUTED, italic=True, after=0)

    # Page 3 - Sections 4 and 5
    new_page(doc, 3, "Week 1 Delivery", "Work Completed During the Week", "The work is organized as a partner-reviewable product definition package, not as a claim of production readiness.")
    add_fact_table(doc, [
        ("Interface", "All navigation, cards, tables, scenarios, modals, notifications, filters, accessible labels, and dynamic workflow copy translated into consistent institutional English."),
        ("Executive journey", "Overview -> Funds -> Fund Detail / Allocation -> Performance -> Internal Comparison -> Peer Comparison, with functional navigation and fund drill-down."),
        ("Analyst journey", "Import -> structure control -> anomaly review -> quarantine -> validation -> local synthetic publication -> lineage, implemented as an interactive simulation."),
        ("Governance cues", "Visible source, run, freshness, quality, hypothesis, quarantine, and lineage indicators; synthetic and sample states remain explicit."),
        ("Verification", "Every page and scenario combination tested; allocation drill-down, modal, keyboard navigation, responsive layouts, and console behavior checked."),
    ])
    add_heading(doc, "Dashboard Navigation and Information Architecture", 2)
    nav_rows = [
        ("Executive", "Overview; Funds; Allocation; Performance; Comparison; Peers", "Move from portfolio signal to fund detail and comparative context."),
        ("Operations", "Data Quality; Import & Validation; Runs & Lineage", "Inspect controls, simulate governed processing, and trace published evidence."),
        ("Global context", "Fund, period, prototype scenario, user journey", "Keep user context visible across sections."),
    ]
    add_matrix(doc, ["Navigation group", "Sections", "Decision supported"], nav_rows, [1900, 3900, 4200], 7.7)
    add_callout(doc, "Important boundary", "No production data connector, authentication, role-based authorization, external publication, or final financial certification is demonstrated. The role selector and publication path are prototype controls.", fill=NAVY_2)

    # Page 4 - Overview
    new_page(doc, 4, "Dashboard Screens", "Executive Overview", "Point-by-point review of the live translated screen.")
    add_screenshot(doc, "01-executive-overview.png", "Figure 1. Executive Overview with navigation, global context, sample-data disclosure, portfolio signal, chart, and KPI cards.", 6.35)
    add_fact_table(doc, [
        ("Objective / user", "Daily portfolio orientation for Executive users."),
        ("Displayed", "Consolidated legal AUM, look-through hypothesis, publication quality, 12-month sample performance, priority funds, and alerts."),
        ("Visuals", "One indexed line chart, four KPI cards, fund table, and alert list."),
        ("Filters / actions", "Fund, period, scenario, journey; View controls and Explore funds."),
        ("Data states", "Sample, current, late, incomplete, loading, empty, error, denied, no match, and hypothesis."),
        ("Data / limits", "Synthetic values from the local prototype Serving payload; no production NAV or AUM certification."),
        ("Product contribution", "Establishes the decision-first landing page and makes trust signals visible beside portfolio metrics."),
        ("Status", "Implemented and visible - screen behavior is verified; business values remain sample data."),
    ], font_size=7.8)

    # Page 5 - Funds and allocation
    new_page(doc, 5, "Dashboard Screens", "Funds, Fund Detail and Allocation", "One fund universe screen and one combined detail/allocation screen are currently present.")
    add_screenshot(doc, "02-fund-allocation-filtered.png", "Figure 2. FUND_01 detail with KPI cards, interactive allocation, reconciliation summary, and a filtered holdings table below the viewport.", 6.35)
    add_fact_table(doc, [
        ("Funds screen", "Lists 14 aliased funds, differentiates history-only from snapshot-plus-history coverage, and supports search and fund selection."),
        ("Detail objective", "Explain one selected fund from net assets and NAV per share through asset allocation and individual positions."),
        ("KPIs", "Net assets, NAV per share, Top 5 concentration, and maximum price age."),
        ("Visuals / tables", "Interactive allocation donut, reconciliation mini-bars, and holdings table with issuer, value, weight, and price freshness."),
        ("Actions", "Select fund; filter holdings by asset class; reset filter; continue to performance or return to funds."),
        ("Drill-down", "Allocation segment -> filtered positions is functional and keyboard-accessible through buttons."),
        ("Limits", "FUND_01 synthetic positions are reused as the demonstration detail baseline; no live position or NAV reconciliation."),
        ("Product contribution", "Defines the portfolio transparency path and the evidence needed for future automated reconciliation."),
        ("Status", "Interactive prototype - navigation and allocation drill-down work with synthetic fixtures."),
    ], font_size=7.45)

    # Page 6 - Performance and comparison
    new_page(doc, 6, "Dashboard Screens", "Performance, Risk and Internal Comparison", "Historical analytics are shown as controlled sample outputs; not all final risk inputs are available.")
    add_screenshot(doc, "03-performance-risk.png", "Figure 3. Performance and Risk view with a base-100 chart, sample metrics, and explicit prerequisite states.", 6.35)
    add_fact_table(doc, [
        ("Performance objective", "Review one fund over the selected period against a synthetic benchmark."),
        ("KPIs / chart", "Cumulative return, volatility, maximum drawdown, Sharpe availability, and one base-100 comparison chart."),
        ("Prerequisites", "Daily NAV is marked controlled; risk-free rate is pending validation; duration is blocked until its unit is agreed."),
        ("Internal Comparison", "Seven snapshot funds are compared at one date by AUM, relative size, daily return, volatility, and quality."),
        ("Filters / actions", "Fund and period selection; links back to Allocation and forward to Comparison / Peers."),
        ("States", "Incomplete and hypothesis states modify metric availability without silently substituting values."),
        ("Limits", "Displayed analytics are synthetic; formula certification, approved reference datasets, and production history are not demonstrated."),
        ("Product contribution", "Defines the final analytical hierarchy while protecting users from unsupported metrics."),
        ("Status", "Partially implemented - views exist; final calculation modules and reference reconciliations remain future work."),
    ], font_size=7.55)

    # Page 7 - Peers
    new_page(doc, 7, "Dashboard Screens", "Peer Comparison", "The interface is deliberately separated from any claim that an external peer universe is certified.")
    add_screenshot(doc, "05-peer-comparison.png", "Figure 4. Synthetic peer universe with activation conditions and explicit sample-data status.", 6.35)
    add_fact_table(doc, [
        ("Objective / user", "Provide comparative context for Executive users when a qualified peer source becomes available."),
        ("Displayed", "Six aliased peers, rank, return index, and sample-data status."),
        ("Controls", "Versioned source, unique universe identifier, and extreme-value quarantine conditions."),
        ("Filters / actions", "Global fund, period, and scenario filters; return to Internal Comparison."),
        ("Data", "Synthetic peer sample only. No real peer fund names or identifiers are exposed."),
        ("Functional interaction", "Page navigation and scenario behavior are implemented; external-source activation is not."),
        ("Limits", "No authorized, stable, qualified peer source or approved taxonomy is connected."),
        ("Product contribution", "Provides a safe adapter concept that can be activated only after source governance is approved."),
        ("Status", "Pending partner validation - source authorization, ownership, taxonomy, reference date, and outlier treatment remain open."),
    ], font_size=7.7)

    # Page 8 - Quality
    new_page(doc, 8, "Dashboard Screens", "Data Quality", "A visible control surface for anomalies, status, severity, and opaque record references.")
    add_screenshot(doc, "06-data-quality.png", "Figure 5. Data Quality view showing synthetic control counts and an anomaly log without source values.", 6.35)
    add_fact_table(doc, [
        ("Objective / user", "Help Analysts determine whether a synthetic batch can proceed and why records are isolated."),
        ("KPIs", "10 open anomalies, 4 blocking anomalies, 49 accepted rows, and 0 confidentiality hits in the current fixture run."),
        ("Table", "Rule, anomaly, severity, status, opaque reference, and Details action."),
        ("Modal", "Displays rule, severity, opaque reference, proposed action, and privacy disclosure; resolution and quarantine actions are simulated."),
        ("Messages", "Source content is excluded from the shared log; blocking and warning states are visually distinct."),
        ("Data", "Generated synthetic quality cases and deterministic run metadata."),
        ("Limits", "The UI demonstrates governance behavior; it does not prove coverage of production source rules or approved tolerances."),
        ("Product contribution", "Makes exceptions actionable and prepares the operational model for coded quality gates."),
        ("Status", "Interactive prototype - anomaly detail and local status updates work on synthetic records."),
    ], font_size=7.55)

    # Page 9 - Import
    new_page(doc, 9, "Dashboard Screens", "Import and Validation", "The Analyst workflow is functional as a simulation and explicitly prevents external publication.")
    add_screenshot(doc, "07-import-control.png", "Figure 6. Analyst workflow after the structure check, with completed steps, anomalies, and actionable details.", 6.35)
    add_fact_table(doc, [
        ("Objective / user", "Guide an Analyst from source selection to controlled local publication."),
        ("Workflow", "Selection -> Structure -> Anomalies -> Treatment -> Validation -> Publication -> Lineage."),
        ("Actions", "Select fixture batch, run structure check, inspect anomaly, quarantine blocking cases, validate, publish locally, and reset."),
        ("Validation", "Silver/Gold contract presence, confidentiality scan, and deterministic fingerprints are presented in the simulated checklist."),
        ("Data", "Three fixture shapes and a logical seven-export batch; classification remains synthetic-example."),
        ("Confirmation", "Toasts confirm structure, dataset validation, and local publication; no external transfer occurs."),
        ("Limits", "No file picker for production sources, human approval service, external queue, or production Serving destination."),
        ("Product contribution", "Defines the safe end-to-end operating sequence required before real ingestion is enabled."),
        ("Status", "Interactive prototype - all seven steps are navigable locally; production ingestion remains unimplemented."),
    ], font_size=7.5)

    # Page 10 - Runs
    new_page(doc, 10, "Dashboard Screens", "Run History and Lineage", "A deterministic evidence path is visible without revealing source values.")
    add_screenshot(doc, "08-runs-lineage.png", "Figure 7. Run log, three screen-to-source lineage proofs, and the beginning of the evidence chain.", 6.35)
    add_fact_table(doc, [
        ("Objective / user", "Allow Analysts and reviewers to inspect processing status and trace representative screen values."),
        ("Run table", "Run identifier, local publication status, generation time, accepted count, quarantined count, and manifest reference."),
        ("Lineage table", "Screen value, KPI, lineage reference, opaque record, and business date."),
        ("Evidence chain", "Manifest -> Silver record -> Gold KPI -> screen value."),
        ("Actions", "Start a new simulated run by returning to Import and Validation."),
        ("Data", "Deterministic synthetic manifest, run, and lineage proofs."),
        ("Limits", "No production orchestration history, job retries, operator identity, retention policy, or source-system audit integration."),
        ("Product contribution", "Establishes the traceability pattern required for governed reporting and future audit evidence."),
        ("Status", "Partially implemented - synthetic run evidence is visible; full automated lineage remains later work."),
    ], font_size=7.55)

    # Page 11 - inventory and states
    new_page(doc, 11, "Dashboard Screens", "Screen Inventory, States and Verification", "All current sections are accounted for; no absent future page is counted as delivered.")
    inventory = [
        ("Executive Overview", "Implemented and visible", "Portfolio signal, KPIs, chart, alerts", "Synthetic portfolio only"),
        ("Funds", "Implemented and visible", "14 aliases, search, selection, coverage", "No production master data"),
        ("Fund Detail / Allocation", "Interactive prototype", "KPI detail, donut, holdings drill-down", "FUND_01 fixture baseline"),
        ("Performance and Risk", "Partially implemented", "Chart, return, volatility, drawdown", "Sharpe and duration dependencies"),
        ("Internal Comparison", "Implemented and visible", "Seven-fund comparable table", "Synthetic snapshot"),
        ("Peer Comparison", "Pending partner validation", "Six aliased sample peers", "No certified source"),
        ("Data Quality", "Interactive prototype", "Anomaly log, modal, status actions", "Synthetic control set"),
        ("Import and Validation", "Interactive prototype", "Seven-step Analyst workflow", "No production ingestion"),
        ("Runs and Lineage", "Partially implemented", "Run and lineage evidence", "No production orchestrator"),
    ]
    add_matrix(doc, ["Current screen", "Status", "Verified evidence", "Current limit"], inventory, [2200, 2200, 3000, 2600], 6.8, status_col=1)
    add_heading(doc, "State and interaction coverage", 2)
    add_fact_table(doc, [
        ("Scenarios", "Sample, current, loading, late, incomplete, empty, error, access denied, no match, and hypothesis pending validation."),
        ("Interactions", "Navigation, fund selection, search, allocation filter, role switch, anomaly modal, status actions, full Analyst workflow, and help notification."),
        ("Viewports", "Desktop baseline and mobile layout verified; intermediate 1024 px and 768 px checks are included in final QA."),
        ("Accessibility", "Skip link, semantic controls, focus transfer, arrow-key sidebar navigation, Escape behavior, labels, and screen-reader-only loading text."),
        ("Technical result", "No console errors; no inline onclick handlers; no horizontal overflow at 390 px."),
    ], font_size=7.7)
    add_status(doc, "Implemented and visible", "The test suite covers all 90 page/scenario combinations.")

    # Page 12 - Executive journey
    new_page(doc, 12, "User Journey", "Executive User Journey", "The journey prioritizes rapid orientation, progressive detail, and visible trust signals.")
    add_callout(doc, "Decision path", "Overview -> prioritize a fund -> inspect allocation and holdings -> review performance -> compare internally -> open peers only when its source conditions are understood.", compact=True)
    steps = [
        ("1. Orient", "Read consolidated AUM, sample performance, publication quality, alerts, source, freshness, and scenario status."),
        ("2. Prioritize", "Open the 14-fund universe and select an aliased fund based on coverage, AUM, daily return, and quality."),
        ("3. Explain", "Inspect NAV per share, concentration, price age, allocation, reconciliation concept, and filtered holdings."),
        ("4. Evaluate", "Review performance and risk indicators, including unavailable prerequisites rather than hidden substitutions."),
        ("5. Compare", "Use internal comparison first; treat the peer view as a synthetic adapter pending source validation."),
    ]
    add_matrix(doc, ["Journey step", "What the current prototype supports"], steps, [2300, 7700], 7.6)
    add_metric_grid(doc, [
        ("14", "Aliased funds", "Seven also have the detailed snapshot demonstration.", PURPLE),
        ("6", "Executive sections", "Overview through peer comparison.", TEAL),
        ("1", "Functional drill-down", "Allocation segment to filtered holdings.", ORANGE),
        ("10", "Scenario states", "Including delayed, empty, error, denied, and hypothesis.", CORAL),
    ], compact=True)
    add_fact_table(doc, [
        ("Implemented now", "Navigation, filters, drill-down, chart/table reading, state communication, and mobile layout."),
        ("Not final behavior", "Exports, authenticated entitlements, live data, certified calculations, and decision logging."),
        ("Partner review", "Confirm priority users, five homepage KPIs, active fund universe, fund-of-funds treatment, and peer scope."),
    ], font_size=7.35)

    # Page 13 - Analyst journey
    new_page(doc, 13, "User Journey", "Analyst User Journey", "The current flow demonstrates safe operating intent with synthetic fixtures and opaque references.")
    workflow_rows = [
        ("1", "Selection", "Choose the authorized synthetic fixture batch.", "Interactive prototype"),
        ("2", "Structure", "Check expected file shapes and record a deterministic outcome.", "Interactive prototype"),
        ("3", "Anomalies", "Inspect rule, severity, opaque record, and proposed action.", "Interactive prototype"),
        ("4", "Treatment", "Quarantine blocking cases without exposing source values.", "Interactive prototype"),
        ("5", "Validation", "Review contract, confidentiality, and idempotence checks.", "Interactive prototype"),
        ("6", "Publication", "Publish only to local synthetic Serving.", "Interactive prototype"),
        ("7", "Lineage", "Trace representative screen values to opaque synthetic records.", "Partially implemented"),
    ]
    add_matrix(doc, ["Step", "Stage", "Current behavior", "Status"], workflow_rows, [700, 1600, 5100, 2600], 7.4, status_col=3)
    add_heading(doc, "Current prototype versus final operating model", 2)
    add_fact_table(doc, [
        ("Source selection", "Now: a fixed synthetic batch. Final: authorized source ingestion with ownership, access, manifest, and environment controls."),
        ("Quality decision", "Now: simulated local status changes. Final: coded rules, approved thresholds, durable quarantine, accountable override, and audit evidence."),
        ("Validation", "Now: prototype checklist. Final: reconciled reference results, sign-off roles, and release criteria."),
        ("Publication", "Now: local synthetic Serving only. Final: governed curated datasets, monitored application services, and role-based access."),
        ("Lineage", "Now: three evidence proofs. Final: automated source-to-screen lineage for every published KPI."),
    ], font_size=7.65)

    # Page 14 - KPIs
    new_page(doc, 14, "Data Display", "KPI, Charts and Data Displayed", "The prototype separates displayed values from the controls and assumptions that qualify them.")
    kpi_rows = [
        ("Portfolio snapshot", "Gross AUM; look-through AUM; publication quality; 12-month performance", "K01, K14, K16, H02", "Cards + overview chart", "Sample / hypothesis"),
        ("Fund detail", "Net assets; NAV per share; Top 5 concentration; maximum price age", "K01, K02, K07, K15", "Cards + allocation donut + holdings", "Sample"),
        ("Performance / risk", "Cumulative return; volatility; maximum drawdown; Sharpe availability", "H02-H05", "Cards + base-100 line chart", "Sample / unavailable"),
        ("Comparison", "AUM; relative size; daily return; volatility; quality", "K01, K03, H03", "Internal table", "Sample"),
        ("Peers", "Rank; aliased peer; return index; source status", "H07", "Peer table", "Synthetic adapter"),
        ("Quality", "Open anomalies; blocking; accepted rows; confidentiality hits", "K16, DQ", "Cards + anomaly table", "Synthetic run"),
        ("Runs / lineage", "Accepted; quarantined; manifest; three trace proofs", "Run + lineage refs", "Two tables + evidence chain", "Synthetic run"),
    ]
    add_matrix(doc, ["Display group", "Measures", "Identifiers", "Visual form", "Current data state"], kpi_rows, [1800, 3000, 1500, 1900, 1800], 6.3)
    add_heading(doc, "Display rules already applied", 2)
    add_bullet(doc, "Every primary screen exposes synthetic/sample status and a stable run reference.", size=8.2, after=1.5)
    add_bullet(doc, "Freshness and quality states remain visible rather than being collapsed into a single success label.", size=8.2, after=1.5)
    add_bullet(doc, "Unavailable prerequisites are named; the prototype does not invent a risk-free rate or duration convention.", size=8.2, after=1.5)
    add_bullet(doc, "Peer values are separated from any certified external universe and use aliases only.", size=8.2, after=1.5)
    add_bullet(doc, "Hard-coded sample figures support UX validation only and must not be interpreted as partner portfolio results.", size=8.2, after=1.5)
    add_callout(doc, "Interpretation rule", "Implemented and visible refers to screen behavior. It does not certify the financial values, formulas, or production data sources.", fill=NAVY_2, compact=True)

    # Page 15 - Governance
    new_page(doc, 15, "Governance", "Data Quality and Governance Elements Already Visible", "Trust is communicated in the product surface while final governance remains a delivery across later sprints.")
    governance_rows = [
        ("Synthetic classification", "Visible across the shell and source notes", "Implemented and visible"),
        ("Source / freshness / quality / run", "Displayed in page footnotes and global status", "Implemented and visible"),
        ("Anomaly severity and status", "Blocking, warning, info, open, quarantined, resolved", "Interactive prototype"),
        ("Opaque record references", "Logs avoid direct source values and private identifiers", "Implemented and visible"),
        ("Quarantine behavior", "Blocking records can be isolated in the simulation", "Interactive prototype"),
        ("Deterministic manifest", "Stable synthetic run metadata and checksums", "Partially implemented"),
        ("Lineage proofs", "Three screen-to-record traces and evidence chain", "Partially implemented"),
        ("Role journey selector", "Executive / Analyst navigation shortcut only", "Interactive prototype"),
        ("LGPD / RBAC controls", "Requirements documented but not enforced by the prototype", "Planned for a later sprint"),
    ]
    add_matrix(doc, ["Element", "Evidence currently visible", "Status"], governance_rows, [3000, 4500, 2500], 7.3, status_col=2)
    add_heading(doc, "Visible evidence chain", 2)
    chain = doc.add_table(rows=1, cols=4)
    set_table_geometry(chain, [2500, 2500, 2500, 2500], 0)
    set_table_borders(chain, LINE, 5)
    for idx, (label, value, fill) in enumerate([
        ("MANIFEST", "SPRINT2-WF-001", LAVENDER),
        ("SILVER", "REC_SNAPSHOT_01", TEAL_LIGHT),
        ("GOLD", "LIN_KPI_K01", ORANGE_LIGHT),
        ("SCREEN", "executive-aum", RED_LIGHT),
    ]):
        shade_cell(chain.cell(0, idx), fill)
        add_text(chain.cell(0, idx), label, 7.3, PURPLE, True, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(chain.cell(0, idx), value, 7.1, INK, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_heading(doc, "Confidentiality statement", 2)
    add_text(doc, "This report uses dashboard screenshots containing only synthetic and aliased data. It includes no real fund names, CNPJ, private ISIN, personal data, local paths, secrets, or private repository identifiers.", 8.6, INK, after=5)

    # Page 16 - Alignment matrix 1
    new_page(doc, 16, "Final Product Alignment", "Alignment with the Final Product - Data and Analytics", "Status is based on executed prototype evidence, not on documentation alone.", landscape=True)
    alignment_1 = [
        ("Governed ETL pipeline", "Partially implemented", "Deterministic local synthetic run and manifest", "Production-grade ingestion, orchestration, error recovery, operating controls", "Sprint 3 v1; Sprint 5 final", "Approved sources and environment"),
        ("Raw/Bronze, Silver, Gold, Serving", "Partially implemented", "Contracts and Silver -> Gold -> Serving evidence", "Persistent governed layers, ownership, retention, deployment", "Sprint 3", "Interface contracts and storage choice"),
        ("Portfolio-source ingestion", "Interactive prototype", "Fixture selection and structure check", "Authorized connectors, secure landing, schema evolution, replay", "Sprint 3", "Source inventory and access"),
        ("NAV and portfolio monitoring", "Partially implemented", "Overview, fund cards, allocation, holdings", "Live curated NAV/positions, reconciliation, alerts", "Sprint 3 alpha", "Fund universe and source mapping"),
        ("Performance indicators", "Partially implemented", "Sample return and base-100 chart", "Certified modules, reference cases, production history", "Sprint 3 alpha", "Formula approval and benchmarks"),
        ("Risk-return analysis", "Partially implemented", "Volatility/drawdown samples; Sharpe unavailable", "Sharpe/Sortino, liquidity, stress, DV01 as prioritized", "Sprint 4", "Risk-free rate, units, thresholds"),
        ("Internal fund comparison", "Interactive prototype", "Seven-fund comparison table", "Certified common-date measures and exclusion rules", "Sprint 3 alpha", "Active universe and validation cases"),
        ("Peer comparison", "Pending partner validation", "Synthetic six-peer adapter view", "Qualified source, taxonomy, outlier rules, approved universe", "Sprint 4 conditional", "Source authorization and availability"),
        ("Automated reconciliation", "Specified but not yet visible", "Reconciliation concept only", "Positions + cash + provisions vs NAV with exceptions", "Sprint 4", "Tolerance and accounting rules"),
        ("Data-quality controls", "Partially implemented", "Synthetic rules, anomaly log, reports", "Production rule coverage, ownership, thresholds, alerts", "Sprint 3-Sprint 4", "Approved domains and tolerances"),
        ("Anomaly quarantine", "Interactive prototype", "Simulated quarantine and opaque log", "Durable quarantine, override workflow, accountable release", "Sprint 3", "Decision rights and retention"),
    ]
    add_matrix(doc, ["Final product capability", "Current Sprint 2 W1 status", "What is currently visible", "What remains to be added", "Target sprint", "Dependency or validation"], alignment_1, [2100, 2100, 2500, 3300, 1700, 2300], 6.0, status_col=1)
    add_text(doc, "Status language follows the seven-level scale defined for this report. Target sprints come from the approved planning sequence; uncertain production decisions are not assigned invented dates.", 7.1, MUTED, italic=True, after=0)

    # Page 17 - Alignment matrix 2
    new_page(doc, 17, "Final Product Alignment", "Alignment with the Final Product - Governance, Operations and Handover", "The role selector and local run evidence are not substitutes for enterprise security or production operations.", landscape=True)
    alignment_2 = [
        ("Lineage and run history", "Partially implemented", "One synthetic run and three trace proofs", "Automated lineage for every output, retries, operator audit", "Sprint 3-Sprint 4", "Orchestrator and metadata model"),
        ("Source freshness", "Interactive prototype", "Current, late, incomplete states and price age", "Source-specific SLAs, monitoring, escalation", "Sprint 3-Sprint 4", "Partner-approved thresholds"),
        ("Authentication", "Specified but not yet visible", "No authentication; demonstration profile only", "Approved identity provider and secure sessions", "Sprint 4", "Hosting and identity decision"),
        ("Role-based access control", "Specified but not yet visible", "Journey selector changes navigation only", "Enforced roles for import, validate, export, and detail", "Sprint 4", "Role matrix approval"),
        ("LGPD and anonymization", "Partially implemented", "Aliases, opaque references, synthetic-only output", "Policy enforcement, access tests, retention, approved export rules", "Sprint 4", "Data classification and security review"),
        ("Production data connections", "Out of scope for Sprint 2 Week 1", "None", "Authorized connectors and environment-specific configuration", "To be confirmed", "Source owners, credentials, hosting"),
        ("Exports and reporting", "Specified but not yet visible", "No PDF/CSV export action", "Filtered, dated, logged, role-controlled outputs", "Sprint 5 / To be confirmed", "Export policy and templates"),
        ("Monitoring and observability", "Specified but not yet visible", "Local run status only", "Health, DQ trends, job alerts, performance telemetry", "Sprint 5", "Deployment architecture"),
        ("User documentation", "Planned for a later sprint", "Prototype labels and this progress report", "Usage guide, known limits, release-aligned screenshots", "Sprint 5; draft Sprint 4", "Stable beta workflows"),
        ("Maintenance documentation", "Planned for a later sprint", "Architecture and operating concepts", "Deployment, replay, monitoring, ownership, troubleshooting", "Sprint 5", "Final architecture and operations"),
        ("Testing, deployment and handover", "Partially implemented", "Automated synthetic tests and responsive QA", "UAT, security, performance, installation, release, knowledge transfer", "Sprint 3-Sprint 5", "Environments, users, acceptance baseline"),
    ]
    add_matrix(doc, ["Final product capability", "Current Sprint 2 W1 status", "What is currently visible", "What remains to be added", "Target sprint", "Dependency or validation"], alignment_2, [2100, 2100, 2500, 3300, 1700, 2300], 5.2, status_col=1)
    add_text(doc, "Authentication, RBAC, production connections, monitoring, exports, manuals, deployment, and final handover are not represented as completed by the prototype.", 7.1, MUTED, italic=True, after=0)

    # Page 18 - Next additions and remaining by sprint
    new_page(doc, 18, "Forward Plan", "Next Additions Not Yet Visible in the Dashboard", "Each horizon distinguishes the current mock-up from the final behavior expected by the roadmap.", landscape=True)
    additions = [
        ("Remainder of Sprint 2", "Validate users, screen baseline, homepage KPIs, source matrix, 7/14-fund treatment, fund-of-funds bridge, peer scope, and Sprint 3 backlog.", "Prevents rework before alpha development.", "Partner review and decision log", "Must", "Pending partner validation", "Today: proposed UI and assumptions. Final S2 outcome: approved or explicitly changed baseline."),
        ("Sprint 3", "Build pipeline v1, governed layers, production-shaped ingestion, coded DQ gates, alpha services, reference tests, and automated lineage.", "Creates a reproducible vertical slice from approved input to screen.", "Approved contracts, sources, KPI cases", "Must", "Planned for a later sprint", "Today: fixtures and local evidence. Final S3 behavior: automated import -> control -> KPI -> alpha screen."),
        ("Sprint 4", "Add advanced risk, reconciliation, controlled peers, RBAC, LGPD controls, audit evidence, end-to-end tests, and UAT.", "Makes the beta analytically useful and safe for pilot validation.", "Stable alpha, identity, qualified peer source or adapter decision", "Must / conditional peers", "Planned for a later sprint", "Today: visual placeholders and samples. Final S4 behavior: governed integrated beta with acceptance evidence."),
        ("Sprint 5", "Stabilize release, monitoring, regression evidence, exports as approved, user/maintenance manuals, deployment, and handover.", "Enables repeatable operation without implicit team knowledge.", "Accepted beta and frozen scope", "Must", "Planned for a later sprint", "Today: working prototype. Final S5 behavior: release package, controlled deployment, manuals, and accepted transfer."),
    ]
    add_matrix(doc, ["Horizon", "Description", "Business reason", "Prerequisite", "Priority", "Current status", "Prototype-to-final difference"], additions, [1200, 2500, 2000, 1900, 900, 1900, 4000], 6.15, status_col=5)
    add_heading(doc, "Remaining Work by Sprint", 2)
    add_fact_table(doc, [
        ("Sprint 2", "Partner decisions, source-of-truth reconciliation, approved prototype baseline, and sprint-ready backlog."),
        ("Sprint 3", "Reproducible data pipeline and working MVP alpha for NAV, portfolio, performance, and internal comparison."),
        ("Sprint 4", "Integrated analytics beta, risk, reconciliation, conditional peers, security, LGPD, and UAT."),
        ("Sprint 5", "Stabilization, release, monitoring, complete documentation, final demonstration, and handover."),
    ], font_size=7.65)
    add_status(doc, "Pending partner validation", "The Sprint 2 exit gate remains open until decisions are recorded.")

    # Page 19 - Dependencies and risks
    new_page(doc, 19, "Control Points", "Dependencies and Pending Partner Decisions", "Open decisions should be resolved before irreversible Sprint 3 implementation choices are made.")
    decisions = [
        ("D1", "Priority users and journeys", "Confirm Executive and Analyst baseline; identify any required Manager or Auditor path.", "YvY Capital"),
        ("D2", "Homepage KPI set", "Approve the five most important metrics and their interpretation conventions.", "YvY Capital + Analytics"),
        ("D3", "Fund universe", "Resolve seven snapshot funds versus fourteen funds with historical coverage.", "YvY Capital + Data"),
        ("D4", "Fund-of-funds bridge", "Approve relationships, look-through treatment, and double-counting controls.", "YvY Capital"),
        ("D5", "Calculation and DQ thresholds", "Approve benchmark, risk-free rate, duration units, freshness, and NAV reconciliation tolerance.", "YvY Capital + Analytics"),
        ("D6", "Peer source", "Confirm authorization, stability, taxonomy, outlier handling, and delivery timing.", "YvY Capital"),
        ("D7", "Access model", "Define who can import, validate, export, and view detailed positions.", "YvY Capital + Security"),
        ("D8", "Environment", "Confirm hosting, identity provider, source access, and demonstration constraints.", "YvY Capital + Inteli"),
    ]
    add_matrix(doc, ["ID", "Decision", "Required outcome", "Decision owner"], decisions, [650, 2350, 5100, 1900], 6.5)
    add_heading(doc, "Risks and Current Limitations", 2)
    risks = [
        ("Late or fragmented review", "Rework or delayed alpha", "Consolidated W4 decision set and bounded corrections", "High"),
        ("7/14-fund ambiguity", "Model and test scope diverge", "Support 14 with a configurable approved subset", "High"),
        ("Unqualified calculations or peers", "Misleading comparison", "Block unsupported metrics; retain explicit sample/pending states", "High"),
        ("Sensitive data exposure", "Confidentiality / LGPD impact", "Authorized private environment, aliases, RBAC, output scanning", "Critical"),
        ("Prototype mistaken for production", "Incorrect business reliance", "Working Draft and synthetic disclosures on report and UI", "High"),
        ("Integration deferred too long", "Unstable beta", "Freeze contracts in Sprint 2 and integrate continuously from Sprint 3", "Medium"),
    ]
    add_matrix(doc, ["Risk / limitation", "Potential impact", "Current control / response", "Level"], risks, [2600, 2500, 4000, 900], 6.5)

    # Page 20 - Next steps and conclusion
    new_page(doc, 20, "Closeout", "Recommended Next Steps", "A short sequence keeps partner decisions early and protects the Sprint 3 critical path.")
    next_steps = [
        "Review the translated prototype with YvY Capital using the Executive and Analyst scenarios shown in this report.",
        "Record D1-D8 as Accepted, Accepted with bounded changes, Not approved, or Deferred, with owner and date.",
        "Freeze the priority screen baseline, fund universe, KPI definitions, source matrix, and peer treatment at the Sprint 2 gate.",
        "Convert approved stories into Sprint 3 work with owner, estimate, fixture, acceptance command, reviewer, and dependency owner.",
        "Begin the Sprint 3 vertical slice with authorized data: ingestion -> controlled layers -> DQ -> KPI -> alpha screen -> lineage evidence.",
        "Keep the prototype and all project evidence in the authorized private environment and preserve synthetic data for demonstrations.",
    ]
    for step in next_steps:
        add_number(doc, step, 8.2, 1.5)
    add_heading(doc, "Conclusion", 2)
    add_text(doc, "Sprint 2 Week 1 has produced a complete English prototype baseline with nine live sections, two coherent user journeys, explicit degraded states, functional drill-down, and visible governance cues. This is meaningful progress toward the Funds Manager product direction.", 8.7, INK, after=3, line=1.06)
    add_text(doc, "The evidence also makes the boundary clear: the current dashboard uses synthetic fixtures and simulated operating logic. Production ingestion, certified financial calculations, automated reconciliation, enterprise access control, monitoring, UAT, deployment, and handover remain later-sprint work.", 8.7, INK, after=3, line=1.06)
    add_callout(doc, "Recommended gate outcome", "Use this report and the translated dashboard as a Working Draft for partner review. Move to Sprint 3 development only after the priority screen set, KPI conventions, source scope, security assumptions, and conditional peer strategy are explicitly recorded.", fill=NAVY, compact=True)
    add_status(doc, "Pending partner validation", "No approval or partner acceptance has been fabricated or implied.")
    add_text(doc, "Prepared on August 22, 2026 | Inteli × YvY Capital | Funds Manager", 7.6, MUTED, italic=True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
